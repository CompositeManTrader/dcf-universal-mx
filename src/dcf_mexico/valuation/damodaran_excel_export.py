"""
Export inputs del modelo al template original Damodaran fcffsimpleginzu.xlsx.

Toma el .xlsx oficial (data/templates/fcffsimpleginzu_template.xlsx) y
sobrescribe las celdas de la hoja "Input sheet" con los valores actuales
del modelo. El usuario descarga un .xlsx completo donde todas las hojas
(Valuation Output, Cost of Capital, Diagnostics, etc.) recalculan
automáticamente con sus inputs porque las fórmulas están preservadas.

Cell map basado en inspección directa del Excel (Jan-2026):
  B3:   Date of valuation
  B4:   Company name
  B7:   Country of incorporation
  B8:   Industry (US)
  B9:   Industry (Global)
  Row 11: Revenues       (B=Most Recent, C=Last 10K, D=Years since 10K)
  Row 12: EBIT
  Row 13: Interest expense
  Row 14: BV equity
  Row 15: BV debt
  B16:  R&D capitalize?  (Yes/No)
  B17:  Op leases?       (Yes/No)
  Row 18: Cash
  B19:  Cross holdings
  B20:  Minority interest
  B21:  Shares outstanding (millones)
  B22:  Stock price
  B23:  Effective tax rate
  B24:  Marginal tax rate
  B26:  Revenue growth Y1
  B27:  Operating margin Y1
  B28:  CAGR Y2-5
  B29:  Target margin
  B30:  Year of convergence
  B31:  S2C Y1-5
  B32:  S2C Y6-10
  B34:  Riskfree
  B37:  Employee options outstanding? (Yes/No)
  B38-B41: option details (count, strike, maturity, stdev)
  B45:  Override CoC after Y10? (Yes/No), B46: CoC value
  B48:  Override ROIC? (Yes/No), B49: ROIC value
  B51:  Override P(failure)? (Yes/No), B52: P value, B53: V/B basis,
        B54: distress %
  B56:  Override reinvest lag? (Yes/No), B57: lag years
  B59:  Override eff→marginal tax convergence? (Yes/No)
  B61:  Override NOL? (Yes/No), B62: NOL value
  B64:  Override rf terminal? (Yes/No), B65: rf terminal
  B67:  Override g terminal? (Yes/No), B68: g terminal
  B70:  Override trapped cash? (Yes/No), B71: trapped cash, B72: tax rate
"""
from io import BytesIO
from pathlib import Path
from typing import Optional
from datetime import date

from openpyxl import load_workbook


def _yn(b: bool) -> str:
    return "Yes" if b else "No"


def _project_root() -> Path:
    return Path(__file__).resolve().parents[3]


def export_to_damodaran_excel(
    base,
    assumptions,
    output,
    *,
    last_10k_revenue: Optional[float] = None,
    last_10k_ebit: Optional[float] = None,
    last_10k_interest: Optional[float] = None,
    last_10k_equity_bv: Optional[float] = None,
    last_10k_debt: Optional[float] = None,
    last_10k_cash: Optional[float] = None,
    years_since_10k: float = 1.0,
    country: str = "Mexico",
    industry_us: str = "Beverage (Alcoholic)",
    industry_global: str = "Beverage (Alcoholic)",
    template_path: Optional[Path] = None,
) -> bytes:
    """Carga el template Damodaran y sobrescribe Input sheet con los valores
    del modelo. Retorna bytes del .xlsx para st.download_button.

    Las fórmulas de las otras hojas (Valuation Output, Cost of capital,
    Diagnostics) recalculan automáticamente al abrir el archivo.
    """
    if template_path is None:
        template_path = _project_root() / "data" / "templates" / "fcffsimpleginzu_template.xlsx"
    if not template_path.exists():
        raise FileNotFoundError(
            f"Template no encontrado en {template_path}. "
            f"Asegúrate de tener fcffsimpleginzu_template.xlsx en data/templates/"
        )

    wb = load_workbook(template_path, keep_vba=False)
    # Encontrar hoja "Input sheet" (case-insensitive)
    target_name = None
    for n in wb.sheetnames:
        if "input" in n.lower():
            target_name = n
            break
    if target_name is None:
        raise ValueError("No se encontró 'Input sheet' en el template")

    ws = wb[target_name]
    a = assumptions
    b = base

    # ============================================================
    # SECCIÓN A · Identification
    # ============================================================
    ws["B3"] = date.today().isoformat()
    ws["B4"] = b.ticker  # Damodaran usa "Company name"; pone el ticker
    ws["B7"] = country
    ws["B8"] = industry_us
    ws["B9"] = industry_global

    # ============================================================
    # SECCIÓN B · Base year numbers (Most Recent 12M | Last 10K | Years)
    # ============================================================
    ws["B11"] = round(b.revenue, 2)
    if last_10k_revenue is not None:
        ws["C11"] = round(last_10k_revenue, 2)
    ws["D11"] = years_since_10k

    ws["B12"] = round(b.ebit, 2)
    if last_10k_ebit is not None:
        ws["C12"] = round(last_10k_ebit, 2)

    ws["B13"] = round(b.interest_expense, 2)
    if last_10k_interest is not None:
        ws["C13"] = round(last_10k_interest, 2)

    ws["B14"] = round(b.equity_book, 2)
    if last_10k_equity_bv is not None:
        ws["C14"] = round(last_10k_equity_bv, 2)

    ws["B15"] = round(b.financial_debt, 2)
    if last_10k_debt is not None:
        ws["C15"] = round(last_10k_debt, 2)

    ws["B16"] = "No"   # R&D capitalize (default for IPC MX)
    ws["B17"] = "No"   # Op leases (IFRS 16 ya capitaliza en MX)

    ws["B18"] = round(b.cash, 2)
    if last_10k_cash is not None:
        ws["C18"] = round(last_10k_cash, 2)

    ws["B19"] = round(b.non_operating_assets, 2)
    ws["B20"] = round(b.minority_interest, 2)

    # ============================================================
    # SECCIÓN C · Market data
    # ============================================================
    # Damodaran espera shares en MILLONES, nuestro base.shares_outstanding
    # está en absolutos.
    ws["B21"] = round(b.shares_outstanding / 1_000_000, 2)
    ws["B22"] = a.market_price or 0
    ws["B23"] = round(a.effective_tax_base, 4)
    ws["B24"] = round(a.marginal_tax_terminal, 4)

    # ============================================================
    # SECCIÓN D · Value drivers
    # ============================================================
    # Y1 growth: si user puso revenue_growth_y1 lo usamos, sino _high
    g_y1 = a.revenue_growth_y1 if a.revenue_growth_y1 is not None else a.revenue_growth_high
    ws["B26"] = round(g_y1, 4)
    # Y1 op margin
    m_y1 = a.op_margin_y1 if a.op_margin_y1 is not None else (
        b.ebit / b.revenue if b.revenue else a.target_op_margin
    )
    ws["B27"] = round(m_y1, 4)
    ws["B28"] = round(a.revenue_growth_high, 4)
    ws["B29"] = round(a.target_op_margin, 4)
    ws["B30"] = a.year_of_margin_convergence
    s2c_15 = (a.sales_to_capital_y1_5
                if a.sales_to_capital_y1_5 is not None
                else a.sales_to_capital)
    s2c_610 = (a.sales_to_capital_y6_10
                  if a.sales_to_capital_y6_10 is not None
                  else a.sales_to_capital)
    ws["B31"] = round(s2c_15, 4)
    ws["B32"] = round(s2c_610, 4)

    # ============================================================
    # SECCIÓN E · Risk inputs (Riskfree es el único editable directo)
    # ============================================================
    ws["B34"] = round(a.risk_free, 4)

    # ============================================================
    # SECCIÓN F · Employee options
    # ============================================================
    n_opts = getattr(a, "options_count", 0.0) or 0.0
    if n_opts > 0:
        ws["B37"] = "Yes"
        ws["B38"] = round(n_opts, 2)
        ws["B39"] = round(getattr(a, "options_strike", 0.0) or 0.0, 2)
        # B40 maturity, B41 stdev — defaults Damodaran
        ws["B40"] = 7
        ws["B41"] = 0.45
    else:
        ws["B37"] = "No"

    # ============================================================
    # SECCIÓN G · Default assumption overrides
    # ============================================================
    # CoC override (B45 Yes/No, B46 value)
    coc_override = a.terminal_wacc_override is not None
    ws["B45"] = _yn(coc_override)
    if coc_override:
        ws["B46"] = round(a.terminal_wacc_override, 4)

    # ROIC terminal override
    ws["B48"] = _yn(a.override_terminal_roic)
    if a.override_terminal_roic:
        ws["B49"] = round(a.terminal_roic_override, 4)

    # Probability of failure
    pf_active = a.probability_of_failure > 0
    ws["B51"] = _yn(pf_active)
    if pf_active:
        ws["B52"] = round(a.probability_of_failure, 4)
        ws["B53"] = a.failure_proceeds_basis  # "V" o "B"
        ws["B54"] = round(a.failure_proceeds_pct, 4)

    # Reinvest lag
    lag_active = a.reinvestment_lag > 0
    ws["B56"] = _yn(lag_active)
    if lag_active:
        ws["B57"] = a.reinvestment_lag

    # NOL
    nol_active = a.nol_carryforward > 0
    ws["B61"] = _yn(nol_active)
    if nol_active:
        ws["B62"] = round(a.nol_carryforward, 2)

    # Terminal rf override
    ws["B64"] = _yn(a.override_terminal_riskfree)
    if a.override_terminal_riskfree:
        ws["B65"] = round(a.terminal_riskfree_override, 4)

    # Terminal growth (default override siempre, ya que damodaran hace
    # que rf default = g default y nuestro modelo lo trata distinto)
    ws["B67"] = "Yes"
    ws["B68"] = round(a.terminal_growth, 4)

    # Trapped cash
    tc_active = a.trapped_cash > 0
    ws["B70"] = _yn(tc_active)
    if tc_active:
        ws["B71"] = round(a.trapped_cash, 2)
        ws["B72"] = round(a.trapped_cash_tax_rate, 4)

    # ============================================================
    # Save to BytesIO and return bytes
    # ============================================================
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
